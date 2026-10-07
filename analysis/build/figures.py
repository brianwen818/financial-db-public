"""Diagrams and charts, drawn into a rectangle of a slide.

Every figure takes ``(slide, x, y, w, h, fs, f)``: a box in inches, a base
font size and the measured facts. The deck places them beside its text; the
figure deck draws one per slide so the same drawing can be exported as an
image for the written report.
"""

from __future__ import annotations

from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

from deck_kit import arrow, box, chip, column_chart, dot, line_chart, node, paint, text
from facts import Facts, n, table, wan

LANES = [("FinMind", "fm"), ("TEJ", "tej")]


def _leading(fs: float, lines: float = 1.0) -> float:
    return fs / 72 * 1.35 * lines


# --------------------------------------------------------------------------
# Diagrams
# --------------------------------------------------------------------------

def architecture(slide, x, y, w, h, fs=13, f: Facts | None = None):
    """Two lanes, four layers: source -> raw -> processed -> views -> use."""
    f = f or Facts()
    lane_w = w * 0.085
    g = w * 0.02
    cw = (w - lane_w - 4 * g) / 5
    head_h = _leading(fs, 2.3)
    gap = _leading(fs, 2.4)
    rh = (h - head_h - gap) / 2
    cols = [x + lane_w + i * (cw + g) for i in range(5)]
    rows = [y + head_h, y + head_h + rh + gap]

    heads = [
        ("來源", "上游系統"),
        ("raw 層", "來源給什麼就存什麼"),
        ("processed 層", "定型・去重・分區"),
        ("view 層", "純 view，可刪掉重建"),
        ("使用", "一律唯讀連線"),
    ]
    for cx, (title, sub) in zip(cols, heads, strict=True):
        text(slide, cx, y, cw, head_h, [
            {"runs": [title], "bold": True, "color": "ink", "size": fs, "after": 0},
            {"runs": [sub], "color": "sub", "size": fs - 2, "after": 0},
        ])

    nodes = {
        "fm": [
            ("FinMind API", ["6 個 endpoint", "＋已下市掃描"]),
            ("原樣 Parquet", ["照上游原樣存", "壞了可以重抓"]),
            ("日分區 Parquet", [f"未還原 {wan(f.fm_raw.n_rows)}", f"還原 {wan(f.fm_adj.n_rows)}"]),
            ("DuckDB view", ["19 個 view", "3 個 macro", "零 table"]),
        ],
        "tej": [
            ("TEJ 增益集", ["住在 Excel 裡", "COM 驅動更新"]),
            ("xlsx 活頁簿", [f"{n(f.tej_adj.securities)}＋48 本", "壞了無法重抓"]),
            ("月分區 Parquet", [f"還原價 {wan(f.tej_adj.n_rows)}", f"指數成分 {wan(f.tej_idx.n_rows)}"]),
            ("DuckDB view", ["20 個 view", "3 個 macro", "零 table"]),
        ],
    }
    for (label, role), ry in zip(LANES, rows, strict=True):
        chip(slide, x, ry + rh / 2 - _leading(fs - 2, 0.7), label, role, size=fs - 2, w=lane_w - 0.1)
        for i, (title, sub) in enumerate(nodes[role]):
            node(slide, cols[i], ry, cw, rh, title, sub, role=role, size=fs)
            arrow(slide, cols[i] + cw + 0.03, ry + rh / 2, cols[i] + cw + g - 0.03, ry + rh / 2)

    node(slide, cols[4], rows[0], cw, 2 * rh + gap, "研究與回測",
         ["Jupyter・SQL・DBeaver", "", "價格、證券主檔、產業、交易日曆、每日指數成分", "",
          "涵蓋率與對帳也是 view，資料健不健康隨時可查"],
         role="ok", size=fs)

    # The one dependency between the pipelines: TEJ seeds missing sessions
    # from FinMind's trading calendar.
    xa, xb = cols[2] + cw / 2, cols[0] + cw / 2
    mid = rows[0] + rh + gap / 2
    arrow(slide, xa, rows[0] + rh, xa, mid, role="sub", dashed=True, head=False)
    arrow(slide, xa, mid, xb, mid, role="sub", dashed=True, head=False)
    arrow(slide, xb, mid, xb, rows[1], role="sub", dashed=True)
    text(slide, xa + 0.12, mid - _leading(fs - 2, 0.5), cols[3] + cw - xa - 0.12, _leading(fs - 2),
         "唯一的相依：TEJ 更新前讀 FinMind 的交易日曆，決定要補哪些交易日", size=fs - 2, color="sub")


def chain(slide, x, y, w, h, steps, fs=13, role="ink", numbered=True):
    """A row of equal nodes joined by arrows. ``steps`` is ``[(title, sub)]``."""
    count = len(steps)
    g = min(0.32, w * 0.03)
    cw = (w - (count - 1) * g) / count
    for i, (title, sub) in enumerate(steps):
        cx = x + i * (cw + g)
        node(slide, cx, y, cw, h, title, sub, role=role, size=fs, number=i + 1 if numbered else None)
        if i < count - 1:
            arrow(slide, cx + cw + 0.03, y + h / 2, cx + cw + g - 0.03, y + h / 2)


def finmind_order(slide, x, y, w, h, fs=13, f=None):
    chain(slide, x, y, w, h, [
        ("維度 ingest", ["主檔、產業、日曆；過濾範圍要用最新名單"]),
        ("價格 ingest", ["回補一檔一次呼叫，每日全市場一次"]),
        ("價格處理", ["改名、定型、過濾、去重、翻成一天一檔"]),
        ("交易日曆", ["1994–98 由價格日期反推（含週六盤）"]),
        ("已下市清單", ["掃描 ∪ 下市公告 − 現有主檔"]),
        ("建立 view", ["重建 view，再跑 28 項健康檢查"]),
    ], fs=fs, role="fm")


def tej_refresh(slide, x, y, w, h, fs=13, f=None):
    """The Excel-driven refresh, ending in a validated save or an abort."""
    main_w = w * 0.74
    chain(slide, x, y, main_w, h, [
        ("讀交易日曆", ["取自 FinMind 的 processed 層"]),
        ("找缺漏日", ["只碰仍在交易的活頁簿"]),
        ("插入 seed 列", ["A 欄代碼、B 欄日期"]),
        ("隱藏 Excel 更新", ["獨立桌面、同步巨集、逾時看門狗"]),
        ("驗證", ["新增列的收盤價是否有值"]),
    ], fs=fs, role="tej")
    g = min(0.32, w * 0.03)
    ox = x + main_w + g
    ow = w - main_w - g
    oh = (h - 0.16) / 2
    node(slide, ox, y, ow, oh, "有值 → 存檔", ["全量重建 Parquet 與 view，再跑 35 項檢查"], role="ok", size=fs,
         pad=0.11)
    node(slide, ox, y + oh + 0.16, ow, oh, "空白 → 不存檔", ["整次中止，不寫半筆資料"], role="ink", size=fs,
         pad=0.11)
    arrow(slide, x + main_w + 0.03, y + h / 2, ox - 0.03, y + oh / 2)
    arrow(slide, x + main_w + 0.03, y + h / 2, ox - 0.03, y + oh + 0.16 + oh / 2)


def schedule(slide, x, y, w, h, fs=12, f=None):
    """One week of scheduled jobs."""
    label_w = w * 0.34
    day_w = (w - label_w) / 7
    head_h = _leading(fs, 1.5)
    jobs = [
        ("FinMind 每日更新", "3 個維度＋缺漏日；秒級", "fm", [0, 1, 2, 3, 4, 5], "18:30", True),
        ("FinMind 每週重抓", "還原價全量＋對帳；約 45 分", "fm", [6], "02:00", True),
        ("TEJ 每日更新", "指數檔＋約 150 本；約 55 分", "tej", [0, 1, 2, 3, 4, 5, 6], "18:45", False),
        ("TEJ 每週全量", "2,453 本全部重抓；約 7 小時", "tej", [6], "03:00", True),
    ]
    legend_h = _leading(fs - 2, 1.8)
    row_h = (h - head_h - legend_h) / len(jobs)
    ly = y + h - _leading(fs - 2, 1.1)
    for lx, label, filled in ((x + label_w, "已註冊的排程", True),
                              (x + label_w + day_w * 2.3, "程式已完成、排程尚未註冊", False)):
        box(slide, lx, ly + 0.02, 0.3, _leading(fs - 2, 0.8), fill="sub" if filled else "white",
            brightness=0.72 if filled else None, line=None if filled else "sub", radius=0.03)
        text(slide, lx + 0.38, ly, day_w * 4.5, _leading(fs - 2), label, size=fs - 2, color="sub")
    for d, label in enumerate("一二三四五六日"):
        text(slide, x + label_w + d * day_w, y, day_w, head_h, f"週{label}", size=fs, color="sub", align="c",
             bold=True)
    for r, (name, cost, role, days, clock, registered) in enumerate(jobs):
        ry = y + head_h + r * row_h
        if r % 2 == 0:
            box(slide, x, ry, w, row_h, fill="tint", radius=0.05)
        dot(slide, x + 0.12, ry + row_h / 2 - _leading(fs, 0.55), 0.12, role)
        text(slide, x + 0.32, ry, label_w - 0.4, row_h, [
            {"runs": [name], "bold": True, "color": "ink", "size": fs, "after": 0},
            {"runs": [cost], "color": "sub", "size": fs - 2, "after": 0},
        ], anchor="m")
        for d in days:
            bw, bh = day_w - 0.12, min(row_h - 0.2, _leading(fs, 1.6))
            shape = box(slide, x + label_w + d * day_w + 0.06, ry + (row_h - bh) / 2, bw, bh,
                        fill=role if registered else "white", brightness=0.72 if registered else None,
                        line=None if registered else role, radius=0.05)
            tf = shape.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            run = p.add_run()
            run.text = clock
            run.font.size = Pt(fs - 1)
            run.font.bold = True
            paint(run.font.color, "ink")


def schema(slide, x, y, w, h, layers, fs=12, role="ink"):
    """Three layer cards with field lists. ``layers`` is ``[(title, sub, lines)]``."""
    g = min(0.36, w * 0.035)
    cw = (w - (len(layers) - 1) * g) / len(layers)
    for i, (title, sub, lines) in enumerate(layers):
        cx = x + i * (cw + g)
        box(slide, cx, y, cw, h, fill="tint")
        dot(slide, cx + 0.18, y + 0.2 + (_leading(fs + 1) - 0.13) / 2 - 0.02, 0.13, role)
        text(slide, cx + 0.4, y + 0.2, cw - 0.58, _leading(fs + 1), title, size=fs + 1, bold=True, color="ink")
        text(slide, cx + 0.18, y + 0.2 + _leading(fs + 1) + 0.04, cw - 0.36, _leading(fs - 1.5, 2), sub,
             size=fs - 1.5, color="sub", spacing=1.05)
        body_y = y + 0.2 + _leading(fs + 1) + _leading(fs - 1.5, 2) + 0.14
        text(slide, cx + 0.18, body_y, cw - 0.36, y + h - body_y - 0.12,
             [{"runs": line if isinstance(line, list) else [line], "after": 3} for line in lines],
             size=fs - 1, spacing=1.08)
        if i < len(layers) - 1:
            arrow(slide, cx + cw + 0.04, y + h / 2, cx + cw + g - 0.04, y + h / 2)


# --------------------------------------------------------------------------
# Charts
# --------------------------------------------------------------------------

def rows_per_year(slide, x, y, w, h, fs=12, f: Facts | None = None):
    """How much of each series exists per year (thousand rows)."""
    df = table("snapshot_rows_per_year")
    # 1994 and the current year are partial years; a part-year total would
    # read as a collapse.
    df = df[(df["year"] >= 1995) & (df["year"] < df["year"].max())]
    cats = [str(v) for v in df["year"]]

    def k(col):
        return [None if v != v else round(v / 1000, 1) for v in df[col]]
    return line_chart(slide, x, y, w, h, cats, [
        ("FinMind 未還原價", k("finmind_unadjusted"), "muted"),
        ("FinMind 還原價", k("finmind_adjusted"), "fm"),
        ("TEJ 還原價", k("tej_adjusted"), "tej"),
    ], fs=fs, number_format="#,##0", tick_skip=5)


def coverage(slide, x, y, w, h, fs=12, f: Facts | None = None):
    """Share of TEJ's exchange-listed rows that FinMind also holds."""
    f = f or Facts()
    df = f.coverage
    cats = [str(v) for v in df["year"]]
    return line_chart(slide, x, y, w, h, cats, [
        ("FinMind 未還原價", list(df["fm_unadjusted_row_pct"]), "muted"),
        ("FinMind 還原價", list(df["fm_adjusted_row_pct"]), "fm"),
    ], fs=fs, number_format='0"%"', y_min=40, y_max=100, tick_skip=5, major_unit=20)


def residual_share(slide, x, y, w, h, market: str, fs=12, f: Facts | None = None, legend=True):
    """Of each year's unadjusted limit breaches, the share still left after adjustment."""
    f = f or Facts()
    df = f.actions
    cats = [str(v) for v in df["year"]]
    base = df[f"{market}_unadjusted_breaches"]
    fm = [round(100 * a / b, 1) if b else None for a, b in zip(df[f"{market}_still_in_fm"], base, strict=True)]
    tej = [round(100 * a / b, 1) if b else None for a, b in zip(df[f"{market}_still_in_tej"], base, strict=True)]
    # TEJ first and FinMind dashed on top: from the year FinMind starts
    # adjusting, the two lines coincide and both must stay visible.
    return line_chart(slide, x, y, w, h, cats, [
        ("TEJ 還原價", tej, "tej"),
        ("FinMind 還原價", fm, "fm", "dash"),
    ], fs=fs, number_format='0"%"', y_min=0, y_max=100, tick_skip=5, major_unit=25, legend=legend)


CAUSE_LABELS = {
    "A": "無成交日的記法不同",
    "B": "本專案管線：還原基準混用",
    "C": "FinMind 未還原除權息",
    "D": "FinMind 多調（未還原價無事件）",
    "E": "兩邊都調、幅度不同",
}


def causes(slide, x, y, w, h, fs=12, f: Facts | None = None):
    """Days on which the two sources' returns differ by more than 1 %, by cause."""
    f = f or Facts()
    keys = ["C", "A", "D", "B", "E"]
    return column_chart(slide, x, y, w, h, [CAUSE_LABELS[k] for k in keys],
                        [("交易日數", [f.cause_days(k) for k in keys], "ink")],
                        fs=fs, labels=True, horizontal=True, legend=False, gap=60,
                        point_roles=["fm", "muted", "fm", "ink", "muted"])


def beyond_limit(slide, x, y, w, h, fs=12, f: Facts | None = None):
    """Returns beyond the daily price limit, on the same rows in each series."""
    f = f or Facts()
    return column_chart(slide, x, y, w, h, ["未還原價", "FinMind 還原價", "TEJ 還原價"],
                        [("筆數", [int(f.limit.unadjusted_beyond_limit), int(f.limit.fm_ticker_beyond_limit),
                                   int(f.limit.tej_beyond_limit)], "ink")],
                        fs=fs, labels=True, legend=False, gap=110, point_roles=["muted", "fm", "tej"])


def case_6669(slide, x, y, w, h, fs=12, f: Facts | None = None):
    """6669 around its 2026-09-02 ex-right date, before the dedupe fix."""
    f = f or Facts()
    df = f.case_6669
    cats = [d[5:] for d in df["date"]]
    return line_chart(slide, x, y, w, h, cats, [
        ("未還原價", list(df["fm_unadjusted"]), "muted"),
        ("FinMind 還原價（修正前的 processed 層）", list(df["fm_processed"]), "fm"),
        ("TEJ 還原價（＝修正後的 FinMind）", list(df["tej"]), "tej"),
    ], fs=fs, number_format="#,##0", tick_skip=3, markers=True, y_min=0)


def zero_price(slide, x, y, w, h, fs=12, f: Facts | None = None):
    """Share of FinMind unadjusted rows whose close is recorded as 0."""
    f = f or Facts()
    df = f.zero[f.zero["year"] >= 2000]
    return column_chart(slide, x, y, w, h, [str(v) for v in df["year"]],
                        [("收盤價為 0 的列占比", list(df["zero_pct"]), "fm")],
                        fs=fs, number_format='0.0"%"', legend=False, tick_skip=5, gap=45)
