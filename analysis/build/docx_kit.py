"""python-docx helpers for the written report: styles, tables, figures, fields."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Microsoft JhengHei"
MONO = "Consolas"
INK = RGBColor(0x17, 0x21, 0x2E)
TEXT = RGBColor(0x11, 0x18, 0x27)
SUB = RGBColor(0x4A, 0x55, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TINT = "F2F4F7"
INK_HEX = "17212E"
LINE_HEX = "D9DEE7"
PAGE_W_CM = 21.0 - 2 * 2.2          # A4 text width

ALIGN = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}


def _fonts(rpr, name: str = FONT) -> None:
    """Latin and East Asian faces must both be set, or CJK falls back to the theme."""
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)


def _style(doc, name, size, bold=False, color=TEXT, before=0, after=6, line=1.25, keep=False):
    style = doc.styles[name]
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    _fonts(style.element.get_or_add_rPr())
    fmt = style.paragraph_format
    fmt.space_before, fmt.space_after = Pt(before), Pt(after)
    fmt.line_spacing = line
    fmt.keep_with_next = keep
    return style


def new_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin, section.bottom_margin = Cm(2.3), Cm(2.2)
    _style(doc, "Normal", 10.5)
    _style(doc, "Heading 1", 17, bold=True, color=INK, before=22, after=8, line=1.2, keep=True)
    _style(doc, "Heading 2", 12.5, bold=True, color=INK, before=14, after=4, line=1.25, keep=True)
    _style(doc, "List Bullet", 10.5, after=3)
    _style(doc, "Caption", 9, color=SUB, before=4, after=12, line=1.25)
    doc.styles["Caption"].font.italic = False
    _page_number_footer(section)
    return doc


def _field(run, instruction: str) -> None:
    for kind, text in (("begin", None), (None, instruction), ("separate", None), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _page_number_footer(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.size = Pt(9)
    run.font.color.rgb = SUB
    _field(run, "PAGE")


def toc(doc) -> None:
    """A table-of-contents field; Word fills it in when fields are updated."""
    p = doc.add_paragraph()
    _field(p.add_run(), 'TOC \\o "1-2" \\h \\z \\u')


def page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _add_runs(p, content, size=None, color=None, bold=False, font=None):
    for item in content if isinstance(content, list) else [content]:
        value, opts = (item, {}) if isinstance(item, str) else item
        run = p.add_run(value)
        run.bold = opts.get("bold", bold)
        if size or opts.get("size"):
            run.font.size = Pt(opts.get("size", size))
        if color is not None or "color" in opts:
            run.font.color.rgb = opts.get("color", color)
        if font or opts.get("mono"):
            name = MONO if opts.get("mono") else font
            run.font.name = name
            _fonts(run._r.get_or_add_rPr(), name)
            if opts.get("mono"):
                # Keep CJK inside a code run in the body face.
                run._r.rPr.find(qn("w:rFonts")).set(qn("w:eastAsia"), FONT)


def para(doc, content, style=None, size=None, color=None, bold=False, align=None, after=None, before=None,
         keep=False):
    p = doc.add_paragraph(style=style)
    _add_runs(p, content, size, color, bold)
    if align:
        p.alignment = ALIGN[align]
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if keep:
        p.paragraph_format.keep_with_next = True
    return p


def heading(doc, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def bullets(doc, items) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        _add_runs(p, item)


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _borders(table, inside: bool = True) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        visible = edge in ("top", "bottom") or (edge == "insideH" and inside)
        el.set(qn("w:val"), "single" if visible else "nil")
        if visible:
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), LINE_HEX)
        borders.append(el)
    tbl_pr.append(borders)


def _cell_margins(cell, top=60, bottom=60, left=100, right=100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def _no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def table(doc, rows, widths, size=9.5, align=None, header=True, bold_first_col=False, mono_cols=()):
    """A table with an ink header row. ``widths`` are centimetres and should sum to the text width."""
    tbl = doc.add_table(rows=len(rows), cols=len(widths))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    _borders(tbl)
    # Word keeps a row with the next one when its paragraphs say so. A short
    # table stays in one piece; a long one only keeps its header attached.
    keep_until = len(rows) - 1 if len(rows) <= 9 else 2
    for r, values in enumerate(rows):
        row = tbl.rows[r]
        _no_split(row)
        is_header = header and r == 0
        if is_header:
            tr_pr = row._tr.get_or_add_trPr()
            tr_pr.append(OxmlElement("w:tblHeader"))
        for c, value in enumerate(values):
            cell = row.cells[c]
            cell.width = Cm(widths[c])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _cell_margins(cell)
            if is_header:
                _shade(cell, INK_HEX)
            for i, line in enumerate(str(value).split("\n")):
                p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.2
                p.paragraph_format.keep_with_next = r < keep_until
                p.alignment = ALIGN[(align or "l" * len(widths))[c]]
                run = p.add_run(line)
                run.font.size = Pt(size)
                run.bold = is_header or (bold_first_col and c == 0)
                if is_header:
                    run.font.color.rgb = WHITE
                elif c in mono_cols:
                    run.font.name = MONO
                    _fonts(run._r.get_or_add_rPr(), MONO)
                    run._r.rPr.find(qn("w:rFonts")).set(qn("w:eastAsia"), FONT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return tbl


def _box(doc, fill: str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.rows[0].cells[0]
    cell.width = Cm(PAGE_W_CM)
    _shade(cell, fill)
    _cell_margins(cell, top=140, bottom=140, left=200, right=200)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl._tbl.tblPr.append(borders)
    return cell


def callout(doc, title: str, lines) -> None:
    """A tinted box: a bold heading and a few paragraphs."""
    cell = _box(doc, TINT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    _add_runs(p, title, size=11, color=INK, bold=True)
    for line in lines:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        _add_runs(p, line)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def code(doc, lines) -> None:
    cell = _box(doc, TINT)
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = i < len(lines) - 1
        run = p.add_run(line)
        run.font.size = Pt(9)
        run.font.name = MONO
        _fonts(run._r.get_or_add_rPr(), MONO)
        run._r.rPr.find(qn("w:rFonts")).set(qn("w:eastAsia"), FONT)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


class Numbering:
    """Running figure and table numbers."""

    def __init__(self) -> None:
        self.figure = 0
        self.table = 0


def figure(doc, counter: Numbering, path: Path, caption: str, width_cm: float = PAGE_W_CM) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    counter.figure += 1
    para(doc, [(f"圖 {counter.figure}　", {"bold": True}), caption], style="Caption", align="c")


def table_caption(doc, counter: Numbering, caption: str) -> None:
    counter.table += 1
    para(doc, [(f"表 {counter.table}　", {"bold": True}), caption], style="Caption", after=4, before=8, keep=True)


def update_fields(path: Path) -> None:
    """Have Word fill in the table of contents and page numbers, then save."""
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    app = win32com.client.DispatchEx("Word.Application")
    app.Visible = False
    app.DisplayAlerts = 0
    try:
        doc = app.Documents.Open(str(path), AddToRecentFiles=False)
        for i in range(1, doc.TablesOfContents.Count + 1):
            doc.TablesOfContents(i).Update()
        doc.Fields.Update()
        doc.Save()
        doc.Close(False)
    finally:
        app.Quit()
