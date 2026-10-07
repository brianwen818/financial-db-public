"""Theme, layouts and drawing helpers shared by the deck and the figure deck.

Colours are theme colours, not literals: a role name maps to a slot of the
theme written by ``new_presentation``, so the deck can be restyled from
PowerPoint's own theme editor. The three accents follow the data source
everywhere -- TEJ is blue, FinMind is orange, and green marks what was
verified or fixed -- so a colour means the same thing on every slide.
"""

from __future__ import annotations

import copy
import re
from xml.sax.saxutils import escape

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE, MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

FONT = "Microsoft JhengHei"

THEME = {
    "dk1": "111827",      # primary text
    "lt1": "FFFFFF",      # surface
    "dk2": "17212E",      # ink: titles, dark slides
    "lt2": "F2F4F7",      # card tint
    "accent1": "2A78D6",  # TEJ
    "accent2": "EB6834",  # FinMind
    "accent3": "1BAF7A",  # verified / fixed / DuckDB
    "accent4": "4A5568",  # secondary text
    "accent5": "8A94A6",  # muted: axes, footers, unadjusted baseline
    "accent6": "D9DEE7",  # hairlines
    "hlink": "2A78D6",
    "folHlink": "4A3AA7",
}

ROLE = {
    "text": MSO_THEME_COLOR.TEXT_1,
    "white": MSO_THEME_COLOR.BACKGROUND_1,
    "ink": MSO_THEME_COLOR.TEXT_2,
    "tint": MSO_THEME_COLOR.BACKGROUND_2,
    "tej": MSO_THEME_COLOR.ACCENT_1,
    "fm": MSO_THEME_COLOR.ACCENT_2,
    "ok": MSO_THEME_COLOR.ACCENT_3,
    "sub": MSO_THEME_COLOR.ACCENT_4,
    "muted": MSO_THEME_COLOR.ACCENT_5,
    "line": MSO_THEME_COLOR.ACCENT_6,
}

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.6
TITLE_Y, TITLE_H = 0.45, 0.9
BODY_Y = 1.55                      # first content row on a titled slide
BODY_W = SLIDE_W - 2 * MARGIN

LAYOUT_TITLE, LAYOUT_SECTION, LAYOUT_CONTENT, LAYOUT_BLANK = 0, 2, 5, 6

_NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def paint(color_format, role: str, brightness: float | None = None) -> None:
    """Set a colour by role. ``brightness`` > 0 gives a tint of it."""
    color_format.theme_color = ROLE[role]
    if brightness is not None:
        color_format.brightness = brightness


# --------------------------------------------------------------------------
# Presentation set-up
# --------------------------------------------------------------------------

def _write_theme(prs) -> None:
    part = prs.slide_master.part.part_related_by(RT.THEME)
    xml = part.blob.decode("utf-8")
    for slot, value in THEME.items():
        xml = re.sub(
            rf"(<a:{slot}>)\s*<a:(?:sysClr|srgbClr)[^>]*/>\s*(</a:{slot}>)",
            rf'\1<a:srgbClr val="{value}"/>\2',
            xml,
        )
    for block in ("majorFont", "minorFont"):
        def fonts(match):
            body = re.sub(r'<a:latin typeface="[^"]*"', f'<a:latin typeface="{FONT}"', match.group(0), count=1)
            body = re.sub(r'<a:ea typeface="[^"]*"', f'<a:ea typeface="{FONT}"', body, count=1)
            # Runs tagged zh-TW resolve through the per-script table, which
            # otherwise maps Traditional Chinese to 新細明體.
            return re.sub(r'(<a:font script="Hant" typeface=")[^"]*"', rf'\1{FONT}"', body)
        xml = re.sub(rf"<a:{block}>.*?</a:{block}>", fonts, xml, flags=re.S)
    xml = re.sub(r'(<a:clrScheme name=")[^"]*"', r'\1financial-db"', xml)
    part._blob = xml.encode("utf-8")


def _lvl1(style_el, size: int, role: str, bold: bool, align: str) -> None:
    """Rewrite the level-1 paragraph defaults of a master text style."""
    lvl = style_el.find(qn("a:lvl1pPr"))
    lvl.set("algn", align)
    rpr = lvl.find(qn("a:defRPr"))
    rpr.set("sz", str(size * 100))
    rpr.set("b", "1" if bold else "0")
    for child in rpr.findall(qn("a:solidFill")):
        rpr.remove(child)
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:schemeClr")).set("val", role)
    rpr.insert(0, fill)


def _place(shape, x, y, w, h) -> None:
    shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(w), Inches(h)


def _placeholder_color(shape, scheme: str, size: int | None = None, align: str | None = None) -> None:
    """Give a layout placeholder its own text colour (overrides the master)."""
    body = shape._element.find(qn("p:txBody"))
    lst = body.find(qn("a:lstStyle"))
    for child in list(lst):
        lst.remove(child)
    lvl = etree.SubElement(lst, qn("a:lvl1pPr"))
    lvl.set("marL", "0")
    lvl.set("indent", "0")
    if align:
        lvl.set("algn", align)
    etree.SubElement(lvl, qn("a:buNone"))
    rpr = etree.SubElement(lvl, qn("a:defRPr"))
    if size:
        rpr.set("sz", str(size * 100))
    fill = etree.SubElement(rpr, qn("a:solidFill"))
    etree.SubElement(fill, qn("a:schemeClr")).set("val", scheme)


def _layout_footer(layout, label: str, w: float, h: float) -> None:
    """Footer label and slide number live on the layout, not on each slide."""
    for ph in list(layout.placeholders):
        if ph.placeholder_format.type is not None and ph.placeholder_format.idx in (10, 11, 12):
            ph._element.getparent().remove(ph._element)
    # python-pptx cannot add shapes to a layout, so the two boxes are written
    # as XML. The slide number is a field, so PowerPoint keeps it current.
    rpr = '<a:solidFill><a:schemeClr val="accent5"/></a:solidFill>'
    label_run = f'<a:r><a:rPr lang="zh-TW" sz="1000">{rpr}</a:rPr><a:t>{escape(label)}</a:t></a:r>'
    number_run = (
        '<a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum">'
        f'<a:rPr lang="zh-TW" sz="1000">{rpr}</a:rPr><a:t>‹#›</a:t></a:fld>'
    )
    tree = layout.shapes._spTree
    for shape_id, name, x, width, align, run in (
        (901, "Footer label", MARGIN, 6.0, "l", label_run),
        (902, "Slide number", w - MARGIN - 1.0, 1.0, "r", number_run),
    ):
        tree.append(etree.fromstring(
            f'<p:sp xmlns:p="{_NS_P}" xmlns:a="{_NS_A}">'
            f'<p:nvSpPr><p:cNvPr id="{shape_id}" name="{name}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{int(Inches(x))}" y="{int(Inches(h - 0.45))}"/>'
            f'<a:ext cx="{int(Inches(width))}" cy="{int(Inches(0.25))}"/></a:xfrm>'
            '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
            '<p:txBody><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"/><a:lstStyle/>'
            f'<a:p><a:pPr algn="{align}"/>{run}</a:p></p:txBody></p:sp>'
        ))


def new_presentation(width: float = SLIDE_W, height: float = SLIDE_H, footer: str | None = None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(width), Inches(height)
    _write_theme(prs)

    master = prs.slide_master
    styles = master._element.find(qn("p:txStyles"))
    _lvl1(styles.find(qn("p:titleStyle")), 30, "tx2", True, "l")
    _lvl1(styles.find(qn("p:bodyStyle")), 16, "tx1", False, "l")
    for ph in master.placeholders:
        if ph.placeholder_format.type is not None and "TITLE" in str(ph.placeholder_format.type):
            _place(ph, MARGIN, TITLE_Y, width - 2 * MARGIN, TITLE_H)
            ph.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    content = prs.slide_layouts[LAYOUT_CONTENT]
    content.name = "Content"
    _place(content.placeholders[0], MARGIN, TITLE_Y, width - 2 * MARGIN, TITLE_H)
    content.placeholders[0].text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    if footer:
        _layout_footer(content, footer, width, height)

    for index, name in ((LAYOUT_TITLE, "Title (dark)"), (LAYOUT_SECTION, "Section (dark)")):
        layout = prs.slide_layouts[index]
        layout.name = name
        layout.background.fill.solid()
        paint(layout.background.fill.fore_color, "ink")
        for ph in list(layout.placeholders):
            if ph.placeholder_format.idx in (10, 11, 12):
                ph._element.getparent().remove(ph._element)
    title_layout = prs.slide_layouts[LAYOUT_TITLE]
    _place(title_layout.placeholders[0], MARGIN + 0.2, 2.2, width - 2 * MARGIN - 0.4, 1.9)
    _placeholder_color(title_layout.placeholders[0], "bg1", 40, "l")
    title_layout.placeholders[0].text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    _place(title_layout.placeholders[1], MARGIN + 0.2, 4.3, width - 2 * MARGIN - 0.4, 1.2)
    _placeholder_color(title_layout.placeholders[1], "accent6", 18, "l")

    section = prs.slide_layouts[LAYOUT_SECTION]
    _place(section.placeholders[0], MARGIN + 0.2, 2.6, width - 2 * MARGIN - 0.4, 1.3)
    _placeholder_color(section.placeholders[0], "bg1", 36, "l")
    section.placeholders[0].text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    _place(section.placeholders[1], MARGIN + 0.2, 4.05, width - 2 * MARGIN - 0.4, 1.2)
    _placeholder_color(section.placeholders[1], "accent6", 16, "l")
    section.placeholders[1].text_frame.vertical_anchor = MSO_ANCHOR.TOP
    return prs


def content_slide(prs, title: str, notes: str | None = None):
    slide = prs.slides.add_slide(prs.slide_layouts[LAYOUT_CONTENT])
    slide.shapes.title.text = title
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[LAYOUT_BLANK])


# --------------------------------------------------------------------------
# Text
# --------------------------------------------------------------------------

def _style_run(run, size, color, bold, italic=False):
    # Marking the run as Chinese turns on PowerPoint's line-breaking rules, so
    # a line never starts with a full stop or comma.
    run.font._element.set("lang", "zh-TW")
    run.font._element.set("altLang", "en-US")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    paint(run.font.color, color)


def _bullet(paragraph, size: float, role: str = "muted") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    indent = int(Pt(size * 1.1))
    ppr.set("marL", str(indent))
    ppr.set("indent", str(-indent))
    clr = etree.SubElement(ppr, qn("a:buClr"))
    scheme = {"muted": "accent5", "tej": "accent1", "fm": "accent2", "ok": "accent3", "ink": "tx2"}[role]
    etree.SubElement(clr, qn("a:schemeClr")).set("val", scheme)
    etree.SubElement(ppr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(ppr, qn("a:buChar")).set("char", "•")


def text(slide, x, y, w, h, content, size=14, color="text", bold=False, align="l", anchor="t",
         spacing=1.12, after=6, name=None):
    """Add a text box.

    ``content`` is a string or a list of paragraphs. A paragraph is a string,
    or a dict ``{"runs": [...], "bullet": bool, "size", "color", "bold",
    "after", "align"}`` whose runs are strings or ``(text, {overrides})``.
    """
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[anchor]
    paragraphs = content if isinstance(content, list) else [content]
    for i, spec in enumerate(paragraphs):
        if isinstance(spec, str):
            spec = {"runs": [spec]}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[spec.get("align", align)]
        p.line_spacing = spacing
        p.space_after = Pt(spec.get("after", after))
        p_size = spec.get("size", size)
        if spec.get("bullet"):
            _bullet(p, p_size, spec.get("bullet_color", "muted"))
        for item in spec["runs"]:
            run_text, opts = (item, {}) if isinstance(item, str) else item
            run = p.add_run()
            run.text = run_text
            _style_run(run, opts.get("size", p_size), opts.get("color", spec.get("color", color)),
                       opts.get("bold", spec.get("bold", bold)), opts.get("italic", False))
    return box


def bullets(items, **kw):
    """Paragraph specs for a bulleted list. An item is a string or a run list."""
    return [{"runs": item if isinstance(item, list) else [item], "bullet": True, **kw} for item in items]


# --------------------------------------------------------------------------
# Shapes
# --------------------------------------------------------------------------

def box(slide, x, y, w, h, fill="tint", brightness=None, line=None, radius=0.06, name=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    if radius:
        # The adjustment is a fraction of the shorter side; keep corners a
        # constant physical size regardless of the box's proportions.
        shape.adjustments[0] = min(0.5, radius / min(w, h))
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        paint(shape.fill.fore_color, fill, brightness)
    if line is None:
        shape.line.fill.background()
    else:
        paint(shape.line.color, line)
        shape.line.width = Pt(1)
    if name:
        shape.name = name
    return shape


def dot(slide, x, y, d, role):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.shadow.inherit = False
    shape.fill.solid()
    paint(shape.fill.fore_color, role)
    shape.line.fill.background()
    return shape


def arrow(slide, x1, y1, x2, y2, role="muted", width=1.5, dashed=False, head=True):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    paint(conn.line.color, role)
    conn.line.width = Pt(width)
    if dashed:
        conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if head:
        ln = conn.line._get_or_add_ln()
        etree.SubElement(ln, qn("a:tailEnd")).set("type", "triangle")
    return conn


def card(slide, x, y, w, h, title, body, role="ink", size=13, title_size=None, fill="tint", pad=0.2):
    """A tinted card: a source-coloured dot, a bold heading and body text."""
    box(slide, x, y, w, h, fill=fill)
    title_size = title_size or size + 2
    d = 0.14
    dot(slide, x + pad, y + pad + (title_size / 72 * 1.2 - d) / 2, d, role)
    text(slide, x + pad + d + 0.1, y + pad, w - 2 * pad - d - 0.1, title_size / 72 * 1.5, title,
         size=title_size, bold=True, color="ink")
    text(slide, x + pad, y + pad + title_size / 72 * 1.5 + 0.08, w - 2 * pad,
         h - 2 * pad - title_size / 72 * 1.5 - 0.08, body, size=size, color="text")


def stat(slide, x, y, w, h, value, label, role="ink", value_size=34, label_size=12, fill="tint"):
    """A stat tile: one large figure with a caption under it."""
    box(slide, x, y, w, h, fill=fill)
    pad = 0.2
    d = 0.12
    dot(slide, x + pad, y + pad, d, role)
    text(slide, x + pad, y + pad + 0.18, w - 2 * pad, value_size / 72 * 1.25, value, size=value_size, bold=True,
         color="ink")
    text(slide, x + pad, y + pad + 0.18 + value_size / 72 * 1.3, w - 2 * pad, h - 2 * pad - 0.18 - value_size / 72 * 1.3,
         label, size=label_size, color="sub")


def chip(slide, x, y, label, role, size=11, w=None):
    """A small pill: tinted fill in the source colour, ink text."""
    w = w or (0.3 + sum(size / 72 * (1.0 if ord(c) > 0x2E7F else 0.56) for c in label))
    h = size / 72 * 1.9
    shape = box(slide, x, y, w, h, fill=role, brightness=0.82, radius=h / 2)
    tf = shape.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = label
    _style_run(run, size, "ink", True)
    return w


def node(slide, x, y, w, h, title, sub=None, role="ink", size=13, fill="white", line="line", sub_size=None,
         number=None, pad=0.14):
    """A diagram node: outlined box, a source-coloured marker, title and sub-lines.

    Title and sub-lines share one text box so a title that wraps pushes the
    rest down instead of overprinting it. With ``number`` the marker is a
    numbered badge on its own row above the text.
    """
    box(slide, x, y, w, h, fill=fill, line=line)
    paragraphs = [{"runs": [title], "bold": True, "color": "ink", "size": size, "after": 3}]
    for line_text in ([] if sub is None else sub if isinstance(sub, list) else [sub]):
        paragraphs.append({"runs": [line_text], "color": "sub", "size": sub_size or size - 2, "after": 1})
    if number is None:
        d = 0.12
        dot(slide, x + pad, y + pad + (size / 72 * 1.3 - d) / 2, d, role)
        text(slide, x + pad + d + 0.08, y + pad, w - 2 * pad - d - 0.08, h - 2 * pad, paragraphs, spacing=1.08)
        return
    d = size / 72 * 1.7
    badge = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + pad), Inches(y + pad), Inches(d), Inches(d))
    badge.shadow.inherit = False
    badge.fill.solid()
    paint(badge.fill.fore_color, role, 0.72)
    badge.line.fill.background()
    tf = badge.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = str(number)
    _style_run(run, size - 1, "ink", True)
    text(slide, x + pad, y + pad + d + 0.07, w - 2 * pad, h - 2 * pad - d - 0.07, paragraphs, spacing=1.08)


# --------------------------------------------------------------------------
# Tables
# --------------------------------------------------------------------------

def table(slide, x, y, widths, rows, size=12, row_h=0.4, header=True, align=None, bold_first_col=False,
          heights=None):
    """A plain table: ink header row, hairline-free body with alternating tint.

    ``rows`` is a list of rows of strings; a cell may be ``(text, role)`` to
    colour its text. ``align`` is one letter per column.
    """
    n_rows, n_cols = len(rows), len(widths)
    heights = heights or [row_h] * n_rows
    frame = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(sum(widths)), Inches(sum(heights)))
    tbl = frame.table
    # Drop the built-in style so only the fills set below show.
    tbl_pr = tbl._tbl.tblPr
    for attr in ("bandRow", "firstRow"):
        tbl_pr.set(attr, "0")
    style = tbl_pr.find(qn("a:tableStyleId"))
    if style is not None:
        style.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
    for c, width in enumerate(widths):
        tbl.columns[c].width = Inches(width)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(heights[r])
        is_header = header and r == 0
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell_text, role = value if isinstance(value, tuple) else (value, None)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if is_header:
                paint(cell.fill.fore_color, "ink")
            else:
                paint(cell.fill.fore_color, "tint" if r % 2 == 0 else "white")
            tf = cell.text_frame
            tf.word_wrap = True
            lines = cell_text.split("\n")
            for i, line_text in enumerate(lines):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                p.alignment = ALIGN[(align or "l" * n_cols)[c]]
                p.line_spacing = 1.05
                run = p.add_run()
                run.text = line_text
                _style_run(run, size, "white" if is_header else (role or "text"),
                           is_header or (bold_first_col and c == 0) or role is not None)
    return frame


# --------------------------------------------------------------------------
# Charts (native, so they stay editable in PowerPoint)
# --------------------------------------------------------------------------

def _style_axes(chart, fs, number_format, y_max, y_min, tick_skip, major_unit):
    cat, val = chart.category_axis, chart.value_axis
    cat.format.line.width = Pt(0.75)
    paint(cat.format.line.color, "muted")
    cat.major_tick_mark = XL_TICK_MARK.NONE
    cat.has_major_gridlines = False
    cat.tick_labels.font.size = Pt(fs)
    paint(cat.tick_labels.font.color, "sub")
    if tick_skip:
        skip = etree.SubElement(cat._element, qn("c:tickLblSkip"))
        skip.set("val", str(tick_skip))
        # Schema order: tickLblSkip sits before tickMarkSkip / noMultiLvlLbl.
        anchor = cat._element.find(qn("c:noMultiLvlLbl"))
        if anchor is not None:
            anchor.addprevious(skip)
    val.format.line.fill.background()
    val.major_tick_mark = XL_TICK_MARK.NONE
    val.has_major_gridlines = True
    val.major_gridlines.format.line.width = Pt(0.5)
    paint(val.major_gridlines.format.line.color, "line")
    val.tick_labels.font.size = Pt(fs)
    paint(val.tick_labels.font.color, "sub")
    val.tick_labels.number_format = number_format
    val.tick_labels.number_format_is_linked = False
    if y_max is not None:
        val.maximum_scale = y_max
    if y_min is not None:
        val.minimum_scale = y_min
    if major_unit is not None:
        val.major_unit = major_unit


def _legend(chart, fs, show):
    chart.has_legend = show
    if show:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(fs)
        paint(chart.legend.font.color, "sub")


def column_chart(slide, x, y, w, h, categories, series, fs=12, number_format="#,##0", legend=None, y_max=None,
                 tick_skip=None, labels=False, gap=70, overlap=-8, horizontal=False, point_roles=None,
                 major_unit=None):
    """Clustered columns (or bars). ``series`` is ``[(name, values, role)]``."""
    data = CategoryChartData()
    data.categories = categories
    for name, values, _ in series:
        data.add_series(name, values)
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.font.size = Pt(fs)
    chart.has_title = False
    plot = chart.plots[0]
    plot.gap_width = gap
    plot.overlap = overlap
    plot.vary_by_categories = False
    for s, (_, _, role) in zip(plot.series, series, strict=True):
        s.format.fill.solid()
        paint(s.format.fill.fore_color, role)
        s.invert_if_negative = False
    if point_roles:
        for i, role in enumerate(point_roles):
            point = plot.series[0].points[i]
            point.format.fill.solid()
            paint(point.format.fill.fore_color, role)
    if labels:
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(fs)
        paint(dl.font.color, "text")
        dl.number_format = number_format
        dl.number_format_is_linked = False
        dl.position = XL_LABEL_POSITION.OUTSIDE_END
    _style_axes(chart, fs, number_format, y_max, None, tick_skip, major_unit)
    if horizontal:
        chart.category_axis.reverse_order = True
        chart.value_axis.visible = not labels
        chart.value_axis.has_major_gridlines = not labels
    _legend(chart, fs, len(series) > 1 if legend is None else legend)
    return chart


def line_chart(slide, x, y, w, h, categories, series, fs=12, number_format="#,##0", legend=None, y_max=None,
               y_min=None, tick_skip=None, markers=False, major_unit=None):
    """Lines. ``series`` is ``[(name, values, role)]`` or with a 4th ``"dash"``."""
    data = CategoryChartData()
    data.categories = categories
    for name, values, *_ in series:
        data.add_series(name, values)
    kind = XL_CHART_TYPE.LINE_MARKERS if markers else XL_CHART_TYPE.LINE
    chart = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data).chart
    chart.font.size = Pt(fs)
    chart.has_title = False
    for s, (_, _, role, *style) in zip(chart.plots[0].series, series, strict=True):
        s.smooth = False
        s.format.line.width = Pt(2.25)
        paint(s.format.line.color, role)
        if style and style[0] == "dash":
            s.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if markers:
            s.marker.style = XL_MARKER_STYLE.CIRCLE
            s.marker.size = 6
            s.marker.format.fill.solid()
            paint(s.marker.format.fill.fore_color, role)
            paint(s.marker.format.line.color, "white")
        else:
            s.marker.style = XL_MARKER_STYLE.NONE
    _style_axes(chart, fs, number_format, y_max, y_min, tick_skip, major_unit)
    _legend(chart, fs, len(series) > 1 if legend is None else legend)
    return chart


def clone_shape(shape, slide):
    """Copy a shape onto another slide (used to reuse a drawn diagram)."""
    el = copy.deepcopy(shape._element)
    slide.shapes._spTree.append(el)
    return el


__all__ = [name for name in dir() if not name.startswith("_")] + ["Emu"]
